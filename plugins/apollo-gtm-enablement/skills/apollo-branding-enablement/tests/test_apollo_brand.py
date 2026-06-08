"""Tests for the apollo_brand toolkit.

Run from the repo root:
    pip install -r plugins/apollo-gtm-enablement/skills/apollo-branding-enablement/tests/requirements.txt
    python -m pytest -q
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

pytest.importorskip("pptx")

SCRIPT_DIR = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

import apollo_brand  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402


# ---- color ----


def test_color_resolves_palette_role():
    c = apollo_brand.color("night")
    assert str(c) == "243031"


def test_color_resolves_raw_hex():
    assert str(apollo_brand.color("1A1A1A")) == "1A1A1A"
    assert str(apollo_brand.color("#1A1A1A")) == "1A1A1A"


def test_palette_has_core_roles():
    for role in ("ink", "white", "paper", "night", "sun", "mist"):
        assert role in apollo_brand.PALETTE


# ---- fonts ----


def test_font_roles_are_scotts_stack():
    assert apollo_brand.FONT_DISPLAY == "Space Grotesk"
    assert apollo_brand.FONT_BODY == "DM Sans"
    assert apollo_brand.FONT_MONO == "DM Mono"


def test_known_fonts_include_stack_and_fallbacks():
    for f in ("Space Grotesk", "DM Sans", "DM Mono", "Arial", "Consolas"):
        assert f in apollo_brand.KNOWN_FONTS


# ---- asset_path ----


def test_asset_path_resolves_existing_logo():
    for variant in ("icon-sun", "icon-stone"):
        p = apollo_brand.asset_path(variant)
        assert p.is_file() and p.name == f"apollo-{variant}.png"


def test_asset_path_raises_for_unknown_variant():
    with pytest.raises(FileNotFoundError):
        apollo_brand.asset_path("icon-purple")


# ---- canvas ----


def test_canvas_matches_reference_decks():
    assert apollo_brand.SLIDE_W == Inches(10.0)
    assert apollo_brand.SLIDE_H == Inches(5.625)


# ---- verify_deck ----


def _new_deck():
    deck = Presentation()
    deck.slide_width = apollo_brand.SLIDE_W
    deck.slide_height = apollo_brand.SLIDE_H
    return deck


def test_verify_deck_clean_on_well_formed_slide(tmp_path):
    deck = _new_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    apollo_brand.set_background(slide, "night")
    apollo_brand.add_text(
        slide, "Hello", left=Inches(0.6), top=Inches(2.0),
        width=Inches(6), height=Inches(1), size=32,
        font=apollo_brand.FONT_DISPLAY, fill="white", autosize=False,
    )
    apollo_brand.add_chrome(slide, bg="night", page_number=1)
    out = tmp_path / "clean.pptx"
    deck.save(str(out))
    assert apollo_brand.verify_deck(str(out)) == []


def test_verify_deck_flags_text_off_slide(tmp_path):
    deck = _new_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    apollo_brand.add_text(
        slide, "way out", left=Inches(9.5), top=Inches(1),
        width=Inches(4), height=Inches(1), size=16, autosize=False,
    )
    out = tmp_path / "offslide.pptx"
    deck.save(str(out))
    issues = apollo_brand.verify_deck(str(out))
    assert any("off-slide" in i for i in issues)


def test_verify_deck_flags_text_collision(tmp_path):
    deck = _new_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    # two same-size boxes drifting into each other (~83% partial overlap)
    apollo_brand.add_text(
        slide, "AAAA", left=Inches(1.0), top=Inches(1),
        width=Inches(3), height=Inches(1), size=16, autosize=False,
    )
    apollo_brand.add_text(
        slide, "BBBB", left=Inches(1.5), top=Inches(1),
        width=Inches(3), height=Inches(1), size=16, autosize=False,
    )
    out = tmp_path / "collide.pptx"
    deck.save(str(out))
    issues = apollo_brand.verify_deck(str(out))
    assert any("collision" in i for i in issues)


def test_verify_deck_flags_unknown_font(tmp_path):
    deck = _new_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    apollo_brand.add_text(
        slide, "stray", left=Inches(1), top=Inches(1),
        width=Inches(3), height=Inches(1), size=16, font="Comic Sans MS", autosize=False,
    )
    out = tmp_path / "font.pptx"
    deck.save(str(out))
    issues = apollo_brand.verify_deck(str(out))
    assert any("unknown font" in i for i in issues)


def test_verify_deck_allows_card_behind_text(tmp_path):
    """A filled card (no text) layered behind a label must NOT be a collision."""
    from pptx.enum.shapes import MSO_SHAPE

    deck = _new_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    card.fill.solid()
    card.fill.fore_color.rgb = apollo_brand.color("mist")
    apollo_brand.add_text(
        slide, "label on card", left=Inches(1.2), top=Inches(1.2),
        width=Inches(3.5), height=Inches(0.4), size=13, autosize=False,
    )
    out = tmp_path / "cardtext.pptx"
    deck.save(str(out))
    assert apollo_brand.verify_deck(str(out)) == []
