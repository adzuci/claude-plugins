"""Starter example for an Apollo GTME deck — card-based, dense, free-form.

This is NOT a renderer to call generically. It is a worked example showing the
reference layout patterns (numbered agenda, two-column content with a sidebar
card, chips, workflow chain) so you can crib the patterns and then DESIGN EACH
SLIDE YOURSELF for the topic at hand. Copy it, rename per topic, rewrite the
slide bodies. Every deck is its own design problem.

Run:
    python3 my_topic_deck.py        # builds + self-verifies

Reference:
- ../SKILL.md       Design reference (palette roles, type scale, slide recipes)
- apollo_brand.py   Light helpers + verify_deck (the post-build guarantee)
- The example decks Scott shared are the visual north star.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apollo_brand import (  # noqa: E402
    FONT_BODY,
    FONT_DISPLAY,
    FONT_MONO,
    SLIDE_H,
    SLIDE_W,
    add_chrome,
    add_speaker_notes,
    add_text,
    color,
    set_background,
    verify_deck,
)

DECK_TITLE = "Session Title Here"
DECK_DATE = "MM.DD.YY"
OUTPUT_PATH = Path("/tmp/gtme-decks/example.pptx")


# ---------------------------------------------------------------------------
# Small layout primitives — cards are the backbone of the reference look
# ---------------------------------------------------------------------------


def card(slide, *, left, top, width, height, fill="mist", radius=0.08, shadow=False):
    """A rounded-rectangle card. Text goes in a SEPARATE add_text on top — the
    card itself carries no text, so it never trips the collision check."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color(fill)
    shp.line.fill.background()
    try:
        shp.adjustments[0] = radius
    except (IndexError, KeyError):
        # shape has no corner-radius adjustment handle — keep its default
        pass
    shp.shadow.inherit = False
    if shadow:
        shp.shadow.inherit = True
    return shp


def chip(slide, text, *, left, top, size=0.42, fill="sun-deep"):
    """A small filled circle chip with a number or glyph centered. `fill` is the
    chip's background color; the centered text is always ink for contrast on the
    light accent fills these chips use."""
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(size), Inches(size))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color(fill)
    shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT_DISPLAY
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = color("ink")
    return shp


# ---------------------------------------------------------------------------
# Slide recipes — author owns every one of these decisions
# ---------------------------------------------------------------------------


def cover(deck, page):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    set_background(slide, "night")
    add_text(
        slide, DECK_TITLE,
        left=Inches(0.6), top=Inches(2.0), width=Inches(8.6), height=Inches(1.1),
        size=48, font=FONT_DISPLAY, fill="white", bold=True, autosize=False,
    )
    add_text(
        slide, DECK_DATE,
        left=Inches(0.62), top=Inches(3.4), width=Inches(3.0), height=Inches(0.4),
        size=14, font=FONT_MONO, fill="sun", autosize=False,
    )
    add_chrome(slide, bg="night", page_number=page)
    add_speaker_notes(slide, "Welcome the team. Note the session is recorded.")


def agenda(deck, page, items):
    """Numbered agenda: left list of numbered rows + right timing sidebar card."""
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    set_background(slide, "paper")
    add_text(
        slide, "Agenda",
        left=Inches(0.6), top=Inches(0.35), width=Inches(6.0), height=Inches(0.6),
        size=32, font=FONT_DISPLAY, fill="ink", bold=True, autosize=False,
    )
    y = 1.3
    for i, (title, sub) in enumerate(items, start=1):
        chip(slide, f"{i:02d}", left=Inches(0.6), top=Inches(y), size=0.40)
        add_text(
            slide, title,
            left=Inches(1.15), top=Inches(y - 0.02), width=Inches(4.4), height=Inches(0.28),
            size=13, font=FONT_DISPLAY, fill="ink", bold=True, autosize=False,
        )
        add_text(
            slide, sub,
            left=Inches(1.15), top=Inches(y + 0.24), width=Inches(4.4), height=Inches(0.24),
            size=10, font=FONT_BODY, fill="muted", autosize=False,
        )
        y += 0.78
    # timing sidebar
    card(slide, left=Inches(6.1), top=Inches(1.1), width=Inches(3.3), height=Inches(3.6), fill="mist")
    add_text(
        slide, "60 MINUTES",
        left=Inches(6.3), top=Inches(1.3), width=Inches(2.9), height=Inches(0.3),
        size=11, font=FONT_MONO, fill="ink", bold=True, autosize=False,
    )
    add_chrome(slide, bg="paper", page_number=page)


def content_two_col(deck, page, *, eyebrow, title, what_label, what_text, side_title, side_rows):
    """Two-column content: left definition card, right labeled rows."""
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    set_background(slide, "paper")
    add_text(
        slide, eyebrow,
        left=Inches(0.6), top=Inches(0.35), width=Inches(4.0), height=Inches(0.26),
        size=11, font=FONT_MONO, fill="ink", bold=True, autosize=False,
    )
    add_text(
        slide, title,
        left=Inches(0.6), top=Inches(0.7), width=Inches(8.8), height=Inches(0.7),
        size=28, font=FONT_DISPLAY, fill="ink", bold=True, autosize=False,
    )
    # left card
    card(slide, left=Inches(0.6), top=Inches(1.7), width=Inches(5.0), height=Inches(2.9), fill="mist", shadow=True)
    add_text(
        slide, what_label,
        left=Inches(0.8), top=Inches(1.85), width=Inches(4.6), height=Inches(0.26),
        size=11, font=FONT_MONO, fill="muted-3", bold=True, autosize=False,
    )
    add_text(
        slide, what_text,
        left=Inches(0.8), top=Inches(2.2), width=Inches(4.6), height=Inches(2.2),
        size=13, font=FONT_BODY, fill="ink",
    )
    # right rows
    add_text(
        slide, side_title,
        left=Inches(5.9), top=Inches(1.7), width=Inches(3.5), height=Inches(0.26),
        size=11, font=FONT_MONO, fill="ink", bold=True, autosize=False,
    )
    y = 2.1
    for label, val in side_rows:
        card(slide, left=Inches(5.9), top=Inches(y), width=Inches(3.5), height=Inches(0.5), fill="mist")
        add_text(
            slide, label,
            left=Inches(6.05), top=Inches(y + 0.08), width=Inches(2.4), height=Inches(0.34),
            size=11, font=FONT_BODY, fill="ink", autosize=False,
        )
        add_text(
            slide, val,
            left=Inches(8.4), top=Inches(y + 0.08), width=Inches(0.9), height=Inches(0.34),
            size=11, font=FONT_MONO, fill="muted-3", align=PP_ALIGN.RIGHT, autosize=False,
        )
        y += 0.62
    add_chrome(slide, bg="paper", page_number=page)


def build() -> Path:
    deck = Presentation()
    deck.slide_width = SLIDE_W
    deck.slide_height = SLIDE_H

    cover(deck, page=1)
    agenda(deck, page=2, items=[
        ("The Field Reality", "Reviewing what's happening IRL"),
        ("What This Is", "Workflow and trigger signals"),
        ("What Good Looks Like", "Gong evidence and patterns"),
        ("Live Practice", "Your accounts, your messaging"),
    ])
    content_two_col(
        deck, page=3,
        eyebrow="THE FIELD REALITY",
        title="What is this intervention?",
        what_label="WHAT IT IS",
        what_text="Replace with the real definition pulled from the Notion source page.",
        side_title="TRIGGER SIGNAL",
        side_rows=[("Signal", "L7"), ("Owner", "AE"), ("SFDC", "Yes")],
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    issues = verify_deck(str(path))
    if issues:
        print(f"verify_deck found {len(issues)} issue(s):")
        for it in issues:
            print(f"  - {it}")
        sys.exit(1)
    print(f"Built deck: {path}")
